"""
Simple Auto-Generate Vietnam Map Excel File
============================================
"""
import os
import sys
import io

# Force UTF-8 encoding for console output
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from pathlib import Path

# Cài đặt thư viện cần thiết
print("Checking and installing required packages...")
os.system('pip install openpyxl python-docx Pillow -q')

import zipfile
import shutil
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font, PatternFill, Alignment

# Đường dẫn
DESKTOP = Path.home() / "Desktop"
WORD_FILE = DESKTOP / "男性_面接者所在地地図2601004.docx"
OUTPUT_DIR = DESKTOP / "VietnamMapAutoMarker"
MAP_IMAGE = OUTPUT_DIR / "BanDoVN.png"
OUTPUT_EXCEL = DESKTOP / "BanDoUngVien.xlsm"

print(f"Desktop: {DESKTOP}")
print(f"Word file exists: {WORD_FILE.exists()}")
print(f"Output: {OUTPUT_EXCEL}")

# Bước 1: Trích xuất ảnh
print("\n[1/3] Extracting map image from Word file...")
if WORD_FILE.exists():
    try:
        temp_dir = OUTPUT_DIR / "temp"
        temp_dir.mkdir(exist_ok=True, parents=True)
        
        with zipfile.ZipFile(WORD_FILE, 'r') as zip_ref:
            zip_ref.extractall(temp_dir)
        
        media_dir = temp_dir / "word" / "media"
        if media_dir.exists():
            images = list(media_dir.glob("*.*"))
            if images:
                largest = max(images, key=lambda x: x.stat().st_size)
                shutil.copy(largest, MAP_IMAGE)
                print(f"   OK: Extracted map image")
            else:
                print(f"   WARNING: No images found in Word file")
        else:
            print(f"   WARNING: Media folder not found")
        
        shutil.rmtree(temp_dir)
    except Exception as e:
        print(f"   ERROR: {e}")
else:
    print(f"   WARNING: Word file not found")

# Bước 2: Tạo Excel
print("\n[2/3] Creating Excel file...")
wb = Workbook()

# Sheet DanhSach
ws1 = wb.active
ws1.title = "DanhSach"
ws1.append(["STT", "Ho va ten", "Tinh/Thanh", "Loai don vi"])

# Format header
for col in range(1, 5):
    cell = ws1.cell(1, col)
    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    cell.font = Font(bold=True, color="FFFFFF")
    cell.alignment = Alignment(horizontal="center", vertical="center")

# Sample data
data = [
    ["01", "NGUYEN DIN TRUNG", "HA NOI", "Thu do"],
    ["02", "PHAM CHUNG HIEU", "HUNG YEN", "Tinh"],
    ["03", "LE PHUOC QUAN", "BINH DINH", "Tinh"],
    ["04", "LE HOANG THACH", "HA TINH", "Tinh"],
    ["05", "TRIEU ANH MINH", "PHU THO", "Tinh"],
    ["06", "LOC HOANG NHAN", "LAM DONG", "Tinh"],
    ["07", "TUONG XUAN LOC", "HA NOI", "Thu do"],
    ["08", "DAM QUANG VINH", "HA NOI", "Thu do"],
    ["09", "TRIEU KHANH HUNG", "PHU THO", "Tinh"],
    ["10", "NGUYEN VAN AN", "HO CHI MINH", "Thanh pho"]
]

for row in data:
    ws1.append(row)

ws1.column_dimensions['A'].width = 8
ws1.column_dimensions['B'].width = 25
ws1.column_dimensions['C'].width = 18
ws1.column_dimensions['D'].width = 12

print(f"   OK: Created sheet 'DanhSach' with {len(data)} candidates")

# Sheet BanDo
ws2 = wb.create_sheet("BanDo")
ws2['A1'] = "Company Name"
ws2['A1'].font = Font(size=12, bold=True)
ws2['A2'] = "Vietnam Map"
ws2['A2'].font = Font(size=14, bold=True)

if MAP_IMAGE.exists():
    try:
        img = ExcelImage(str(MAP_IMAGE))
        img.width = 600
        img.height = int(600 * (img.height / img.width))
        ws2.add_image(img, 'C3')
        print(f"   OK: Added map image to sheet 'BanDo'")
    except Exception as e:
        print(f"   ERROR: Could not add image - {e}")
else:
    print(f"   WARNING: Map image not found at {MAP_IMAGE}")
    print(f"   The Excel file will be created without the map image.")

ws2.column_dimensions['A'].width = 30

# Save
try:
    wb.save(OUTPUT_EXCEL)
    print(f"   OK: Saved Excel file")
except Exception as e:
    print(f"   ERROR: Could not save Excel - {e}")

# Bước 3: Tạo hướng dẫn
print("\n[3/3] Creating instruction file...")
instruction = DESKTOP / "HUONG_DAN_THEM_VBA.txt"
try:
    with open(instruction, 'w', encoding='utf-8') as f:
        f.write("="*60 + "\n")
        f.write("   HUONG DAN THEM VBA VAO FILE EXCEL\n")
        f.write("="*60 + "\n\n")
        f.write(f"File Excel created: BanDoUngVien.xlsm\n\n")
        f.write("NEXT STEPS:\n\n")
        f.write("1. Open file BanDoUngVien.xlsm\n")
        f.write("2. Press Alt + F11 (open VBA Editor)\n")
        f.write("3. Menu: Insert -> Module\n")
        f.write("4. Open file MapGenerator.bas in VietnamMapAutoMarker folder\n")
        f.write("5. Copy all content and paste into Module\n")
        f.write("6. Close VBA Editor (Alt + Q)\n")
        f.write("7. Go to 'BanDo' sheet\n")
        f.write("8. Tab Developer -> Insert -> Button (Form Control)\n")
        f.write("9. Draw button and select macro: GenerateVietnamMap\n")
        f.write("10. Name button: TAO BAN DO\n\n")
        f.write("IMPORTANT: Rename image shape to 'BanDoVN'\n")
        f.write("- Click the map image\n")
        f.write("- In Name Box (top-left), type: BanDoVN\n")
        f.write("- Press Enter\n\n")
        f.write("DONE!\n")
    print(f"   OK: Created instruction file")
except Exception as e:
    print(f"   ERROR: {e}")

print("\n" + "="*60)
print("   COMPLETED!")
print("="*60)
print(f"\nFiles created:")
print(f"  - {OUTPUT_EXCEL.name}")
if MAP_IMAGE.exists():
    print(f"  - {MAP_IMAGE.name}")
print(f"  - {instruction.name}")
print("\nNext: Follow instructions in HUONG_DAN_THEM_VBA.txt")
print("="*60)
